"""Generate ScrapeLite_intro.pptx — a short, non-technical intro deck for peers."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x10, 0x1C, 0x2E)
BLUE = RGBColor(0x25, 0x63, 0xEB)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
INK = RGBColor(0x1C, 0x27, 0x33)
MUTED = RGBColor(0x5B, 0x66, 0x74)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_text(slide, x, y, w, h, runs, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is a list of (text, bold, color) or a str."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, bold, color)]
        for (txt, b, c) in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.color.rgb = c
            r.font.name = "Calibri"
    return tb


def title_slide():
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, NAVY)
    add_rect(s, Inches(0.9), Inches(3.05), Inches(1.6), Inches(0.09), BLUE)
    add_text(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(1.2),
             ["ScrapeLite"], size=60, color=WHITE, bold=True)
    add_text(s, Inches(0.9), Inches(3.35), Inches(11.5), Inches(0.9),
             ["An automatic data scraper for scientific papers"], size=28, color=RGBColor(0xB8, 0xC4, 0xD4))
    add_text(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(0.7),
             ["Search papers  →  pull out the data you need  →  get an Excel file"],
             size=20, color=RGBColor(0x8F, 0xA5, 0xC0))


def content_slide(title):
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.18), BLUE)
    add_text(s, Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.8),
             [title], size=34, color=NAVY, bold=True)
    return s


def bullets(slide, items, x=0.9, y=1.7, w=11.5, h=5.3, size=22):
    runs = []
    for it in items:
        if isinstance(it, tuple):
            lead, rest = it
            runs.append([(lead, True, BLUE), (rest, False, INK)])
        else:
            runs.append([("•  ", True, BLUE), (it, False, INK)])
    add_text(slide, Inches(x), Inches(y), Inches(w), Inches(h), runs,
             size=size, line_spacing=1.15)


# ---- Slide 1: title ----
title_slide()

# ---- Slide 2: the problem ----
s = content_slide("Reading papers by hand is slow")
bullets(s, [
    "You need facts from dozens of papers — mutations, results, numbers, conditions.",
    "Today that means opening each PDF and copying the values into a spreadsheet by hand.",
    "It takes hours and it is easy to miss or mistype things.",
    "We want the data, not the busywork.",
], y=2.1, size=24)

# ---- Slide 3: how it works (3 steps) ----
s = content_slide("What it does, in three steps")
steps = [
    ("1  Search   ", "Type your keywords and pick the databases (PubMed, Europe PMC, Crossref)."),
    ("2  Collect   ", "It finds the papers and automatically downloads the full text of open-access ones."),
    ("3  Extract   ", "An AI reads each paper and fills in exactly the columns you asked for."),
]
add_text(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(4.6),
         [ [ (k, True, BLUE), (v, False, INK) ] for k, v in steps ],
         size=26, line_spacing=1.6)

# ---- Slide 4: what you get ----
s = content_slide("What you get: a clean Excel file")
bullets(s, [
    ("search_results.xlsx   —   ", "every paper it found (title, authors, DOI, link)."),
    ("extracted_data.xlsx   —   ", "the data you asked for, one row per finding."),
    ("closed_source_papers.xlsx   —   ", "paywalled papers it could not open (logged, not skipped)."),
    ("Everything lands in the ", "output"),
], y=2.0, size=24)

# ---- Slide 5: how to run it ----
s = content_slide("Using it takes no code at all")
steps = [
    ("First time only   ", "double-click setup.bat (installs everything once)."),
    ("Every time   ", "double-click start.bat — a page opens in your browser."),
    ("Type your terms, pick sites, click ", "Search"),
    ("Describe what to extract, click ", "Extract"),
    ("Open the ", "output"),
]
add_text(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(4.8),
         [ [ (k, True, GREEN), (v, False, INK) ] for k, v in steps ],
         size=24, line_spacing=1.6)

# ---- Slide 6: why it matters ----
s = content_slide("Why this helps")
bullets(s, [
    "Hours of manual copying become a few minutes of waiting.",
    "Consistent: every paper is read the same way, so nothing gets missed differently.",
    "Private: the AI runs on your own computer — your data never leaves the machine.",
    "You can stop and continue anytime; it remembers what is already done.",
], y=2.1, size=24)

prs.save("ScrapeLite_intro.pptx")
print("Saved ScrapeLite_intro.pptx with", len(prs.slides._sldIdLst), "slides")
