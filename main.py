"""
main.py — run the whole pipeline, or just one stage.

    python main.py                 search, then fetch + extract everything
    python main.py --search-only   only search -> output/search_results.xlsx
    python main.py --extract-only  only fetch + extract (reads search_results.xlsx)
    python main.py --limit 5       extract at most 5 papers (quick demo)
    python main.py --reset         clear the resume checkpoint, then run

See README.md for the full picture.
"""

import argparse
import sys

# Windows consoles default to cp1252, which can't print Greek/math symbols
# that appear in paper titles (α, β, γ, ...). Print as UTF-8 instead of crashing.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import load_workbook

import config
import extractor
import fetcher
import searchers
import writer
from checkpoint import Checkpoint

SEARCH_COLS = [
    "source", "title", "authors", "doi", "pmid", "pmcid",
    "year", "journal", "is_open_access", "fulltext_url",
]
CLOSED_COLS = ["title", "authors", "doi", "source", "reason"]
EXTRACT_COLS = ["doi", "source"] + [e[0] for e in config.FIELDS]


def run_search():
    print("== Step 1/2: searching ==")
    new_papers = searchers.search_all()

    # Load any papers already in search_results.xlsx so a fresh search ADDS
    # to them (deduped) instead of wiping previous rows. If you want a totally
    # different topic, just delete output/search_results.xlsx first — a fresh
    # one is created on the next run.
    existing = _load_search_results(required=False)
    papers = searchers.merge_papers(existing, new_papers)
    added = len(papers) - len(existing)

    rows = [{
        "source": ", ".join(p.get("sources", [])),
        "title": p.get("title", ""),
        "authors": p.get("authors", ""),
        "doi": p.get("doi", ""),
        "pmid": p.get("pmid", ""),
        "pmcid": p.get("pmcid", ""),
        "year": p.get("year", ""),
        "journal": p.get("journal", ""),
        "is_open_access": p.get("is_oa"),
        "fulltext_url": p.get("fulltext_url", ""),
    } for p in papers]
    writer.write_xlsx(config.SEARCH_RESULTS_XLSX, SEARCH_COLS, rows)
    print(f"Wrote {len(rows)} papers ({added} new) to {config.SEARCH_RESULTS_XLSX}")
    return papers


def _load_search_results(required=True):
    if not config.SEARCH_RESULTS_XLSX.exists():
        if required:
            sys.exit(
                f"No {config.SEARCH_RESULTS_XLSX.name} yet. Run `python main.py --search-only` "
                "first (or run `python main.py` to do everything)."
            )
        return []
    wb = load_workbook(config.SEARCH_RESULTS_XLSX, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    papers = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        papers.append({
            "sources": [s.strip() for s in str(d.get("source") or "").split(",") if s.strip()],
            "title": d.get("title") or "",
            "authors": d.get("authors") or "",
            "doi": d.get("doi") or "",
            "pmid": d.get("pmid") or "",
            "pmcid": d.get("pmcid") or "",
            "year": d.get("year") or "",
            "journal": d.get("journal") or "",
            "abstract": d.get("abstract") or "",
            "is_oa": d.get("is_open_access") == "yes",
            "fulltext_url": d.get("fulltext_url") or "",
        })
    return papers


def run_extract(limit=None, papers=None):
    print("== Step 2/2: fetching + extracting ==")
    extractor.check_backend()
    cp = Checkpoint(config.CHECKPOINT_JSON)

    if papers is None:
        papers = _load_search_results()

    # dedup + skip already-done
    todo, seen = [], set()
    for p in papers:
        k = searchers.paper_key(p)
        if k in seen:
            continue
        seen.add(k)
        if cp.is_done(k):
            continue
        todo.append(p)
    if limit:
        todo = todo[:limit]

    print(f"{len(todo)} paper(s) to process.")
    for p in todo:
        k = searchers.paper_key(p)
        title = (p.get("title") or "")[:60]
        print(f"\n[{k}] {title}")

        try:
            text, source = fetcher.fetch_full_text(p)
        except Exception as e:
            cp.mark(k, "error", error=str(e))
            print(f"  !! fetch error: {e}")
            continue

        if not text:
            writer.append_xlsx(config.CLOSED_SOURCE_XLSX, CLOSED_COLS, [{
                "title": p.get("title", ""),
                "authors": p.get("authors", ""),
                "doi": p.get("doi", ""),
                "source": ", ".join(p.get("sources", [])),
                "reason": source,
            }])
            cp.mark(k, "closed", reason=source)
            print(f"  -> closed source: {source}")
            continue

        print(f"  -> got full text ({len(text)} chars) via {source}; extracting ...")
        try:
            rows = extractor.extract(text)
        except Exception as e:
            cp.mark(k, "error", error=str(e))
            print(f"  !! extract error: {e}")
            continue

        rows = [r for r in rows if isinstance(r, dict)]
        field_names = [e[0] for e in config.FIELDS]
        out_rows = [
            {"doi": p.get("doi", ""), "source": ", ".join(p.get("sources", []))}
            | {f: r.get(f, "") for f in field_names}
            for r in rows
        ]
        writer.append_xlsx(config.EXTRACTED_XLSX, EXTRACT_COLS, out_rows)
        cp.mark(k, "extracted", n=len(out_rows))
        print(f"  -> wrote {len(out_rows)} row(s)")

    print(f"\nDone. See {config.EXTRACTED_XLSX} and {config.CLOSED_SOURCE_XLSX}.")


def parse_args():
    ap = argparse.ArgumentParser(description="Generalized search -> extract pipeline.")
    ap.add_argument("--search-only", action="store_true", help="only search (skip extraction)")
    ap.add_argument("--extract-only", action="store_true", help="only fetch + extract from existing search results")
    ap.add_argument("--limit", type=int, default=None, help="max papers to extract (quick demo)")
    ap.add_argument("--reset", action="store_true", help="clear the resume checkpoint before running")
    return ap.parse_args()


def main():
    args = parse_args()
    if args.reset:
        Checkpoint(config.CHECKPOINT_JSON).reset()
        print("Checkpoint cleared.")

    if args.extract_only:
        run_extract(limit=args.limit)
    elif args.search_only:
        run_search()
    else:
        papers = run_search()
        run_extract(limit=args.limit, papers=papers)


if __name__ == "__main__":
    main()
