"""
app.py — local web app for the search -> fetch -> extract pipeline.

Run it with:

    python app.py        (or double-click start.bat)

A browser opens at http://127.0.0.1:5000. No terminal needed after that.

This is a thin Flask layer over the same engine the CLI uses (searchers, fetcher,
extractor, writer, checkpoint), so both interfaces stay in sync.
"""

import json
import os
import subprocess
import sys
import threading
import uuid
import webbrowser
from pathlib import Path

from flask import Flask, abort, jsonify, render_template, request, send_file
from openpyxl import load_workbook

import config
import extractor
import fetcher
import searchers
import writer
from checkpoint import Checkpoint
from main import CLOSED_COLS, SEARCH_COLS, _load_search_results

app = Flask(__name__)

# --------------------------------------------------------------------------
# Settings (persisted server-side so users don't retype email/model each time)
# --------------------------------------------------------------------------

SETTINGS_PATH = config.OUTPUT_DIR / "web_settings.json"


def _default_settings():
    return {
        "email": config.NCBI_EMAIL,
        "model": config.OLLAMA_MODEL,
        "max_results": config.DEFAULT_MAX_RESULTS,
    }


def load_settings():
    if SETTINGS_PATH.exists():
        try:
            data = json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
            d = _default_settings()
            d.update({k: v for k, v in data.items() if k in d})
            return d
        except Exception:
            return _default_settings()
    return _default_settings()


def save_settings(s):
    SETTINGS_PATH.write_text(json.dumps(s, indent=2, ensure_ascii=False), encoding="utf-8")


def apply_settings(s):
    config.NCBI_EMAIL = s.get("email") or config.NCBI_EMAIL
    config.OLLAMA_MODEL = s.get("model") or config.OLLAMA_MODEL
    config.HTTP_HEADERS = {"User-Agent": f"general-datascraper/1.0 (mailto:{config.NCBI_EMAIL})"}


apply_settings(load_settings())

# --------------------------------------------------------------------------
# Background jobs (one at a time)
# --------------------------------------------------------------------------

JOBS = {}
JOB_LOCK = threading.Lock()


def _new_job(job_type):
    job = {
        "id": uuid.uuid4().hex[:8],
        "type": job_type,
        "status": "running",
        "error": None,
        "progress": {"done": 0, "total": 0, "current": ""},
        "log": [],
        "result": None,
    }
    JOBS[job["id"]] = job
    return job


def _active_job():
    for j in list(JOBS.values()):
        if j["status"] == "running":
            return j
    return None


def _start_job(job_type, fn, *args):
    with JOB_LOCK:
        if _active_job():
            return None, "A task is already running. Wait for it to finish, then try again."
        job = _new_job(job_type)
    threading.Thread(target=fn, args=(job, *args), daemon=True).start()
    return job, None


def _job_public(job):
    if not job:
        return None
    return {
        "id": job["id"],
        "type": job["type"],
        "status": job["status"],
        "error": job["error"],
        "progress": job["progress"],
        "log": job["log"][-40:],
        "result": job["result"],
    }


# --------------------------------------------------------------------------
# Job bodies (mirror main.py, but report progress into the job dict)
# --------------------------------------------------------------------------

def _paper_to_row(p):
    return {
        "source": ", ".join(p.get("sources", [])),
        "title": p.get("title", ""),
        "authors": p.get("authors", ""),
        "doi": p.get("doi", ""),
        "pmid": p.get("pmid", ""),
        "pmcid": p.get("pmcid", ""),
        "year": p.get("year", ""),
        "journal": p.get("journal", ""),
        "is_open_access": p.get("is_oa"),
        "fulltext_url": p.get("fulltext_url", ""),
    }


def _run_search_job(job, query, sources, max_results):
    try:
        apply_settings(load_settings())
        job["log"].append(f"Query: {query}")
        job["log"].append(f"Websites: {', '.join(sources)}")
        # A new search starts a fresh dataset: replace the search sheet AND clear
        # any previous extraction outputs + resume state, so Extract only ever
        # works on this search's papers.
        papers = searchers.search(query, sources, limit=max_results)
        rows = [_paper_to_row(p) for p in papers]
        writer.write_xlsx(config.SEARCH_RESULTS_XLSX, SEARCH_COLS, rows)
        for path in (config.EXTRACTED_XLSX, config.CLOSED_SOURCE_XLSX):
            if Path(path).exists():
                Path(path).unlink()
        Checkpoint(config.CHECKPOINT_JSON).reset()
        job["result"] = {"papers": rows, "added": len(rows), "total": len(rows)}
        job["status"] = "done"
        job["log"].append(f"Done: {len(rows)} papers in the sheet (previous data cleared).")
    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)
        job["log"].append(f"Error: {e}")


def make_instruction(goal):
    goal = (goal or "").strip() or "Extract the relevant data described in the text."
    return (
        "You are an expert research assistant. Read the full paper TEXT below very carefully.\n"
        f"Task: {goal}\n\n"
        "Rules:\n"
        "- Extract EVERY record that matches the task. Do not skip any, and do not merge\n"
        "  separate records into one row.\n"
        "- Only report information that is explicitly stated in the TEXT. Do not invent,\n"
        "  infer, or copy values from memory or from outside the paper.\n"
        "- If a value is not stated in the text, leave that field empty instead of guessing.\n"
        "- Keep wording and values consistent with how they appear in the paper.\n"
        "- If no record matches the task, return an empty list []."
    )


def _read_header(path):
    p = Path(path)
    if not p.exists():
        return None
    try:
        wb = load_workbook(p, read_only=True)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        wb.close()
        return header
    except Exception:
        return None


def _run_extract_job(job, goal, columns):
    try:
        apply_settings(load_settings())
        columns = [c.strip() for c in columns if c and str(c).strip()]
        fields = [(c, "text") for c in columns] if columns else list(config.FIELDS)
        field_names = [f[0] for f in fields]
        instruction = make_instruction(goal)
        extract_cols = ["doi", "source"] + field_names

        papers = _load_search_results(required=False)
        if not papers:
            job["status"] = "done"
            job["result"] = {"extracted": 0, "closed": 0, "papers_processed": 0}
            job["log"].append("No search results yet — run a search first.")
            return

        cp = Checkpoint(config.CHECKPOINT_JSON)
        todo, seen = [], set()
        for p in papers:
            k = searchers.paper_key(p)
            if k in seen:
                continue
            seen.add(k)
            if cp.is_done(k):
                continue
            todo.append(p)

        job["progress"]["total"] = len(todo)
        job["log"].append(f"{len(todo)} paper(s) to process (already-done skipped).")

        n_extracted = n_closed = 0
        first_write = True
        for i, p in enumerate(todo):
            k = searchers.paper_key(p)
            title = (p.get("title") or "")[:60]
            job["progress"]["done"] = i
            job["progress"]["current"] = title
            job["log"].append(f"[{i + 1}/{len(todo)}] {title}")

            try:
                text, source = fetcher.fetch_full_text(p)
            except Exception as e:
                cp.mark(k, "error", error=str(e))
                job["log"].append(f"   !! fetch error: {e}")
                continue
            if not text:
                writer.append_xlsx(config.CLOSED_SOURCE_XLSX, CLOSED_COLS, [{
                    "title": p.get("title", ""),
                    "authors": p.get("authors", ""),
                    "doi": p.get("doi", ""),
                    "source": ", ".join(p.get("sources", [])),
                    "reason": source,
                }])
                cp.mark(k, "closed", reason=source)
                n_closed += 1
                job["log"].append(f"   -> closed source ({source})")
                continue

            try:
                rows = extractor.extract(text, fields=fields, instruction=instruction)
            except Exception as e:
                cp.mark(k, "error", error=str(e))
                job["log"].append(f"   !! extract error: {e}")
                continue

            rows = [r for r in rows if isinstance(r, dict)]
            out_rows = [
                {"doi": p.get("doi", ""), "source": ", ".join(p.get("sources", []))}
                | {f: r.get(f, "") for f in field_names}
                for r in rows
            ]
            if out_rows:
                if first_write:
                    existing_header = _read_header(config.EXTRACTED_XLSX)
                    if existing_header == extract_cols:
                        writer.append_xlsx(config.EXTRACTED_XLSX, extract_cols, out_rows)
                    else:
                        writer.write_xlsx(config.EXTRACTED_XLSX, extract_cols, out_rows)
                        if existing_header:
                            job["log"].append("   (columns changed — started a fresh extracted_data.xlsx)")
                    first_write = False
                else:
                    writer.append_xlsx(config.EXTRACTED_XLSX, extract_cols, out_rows)
                n_extracted += len(out_rows)
            cp.mark(k, "extracted", n=len(out_rows))
            job["log"].append(f"   -> {len(out_rows)} row(s)")

        job["progress"]["done"] = len(todo)
        job["progress"]["current"] = ""
        job["result"] = {"extracted": n_extracted, "closed": n_closed, "papers_processed": len(todo)}
        job["status"] = "done"
        job["log"].append(f"Done: {n_extracted} row(s) extracted, {n_closed} paper(s) closed.")
    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)
        job["log"].append(f"Error: {e}")


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/status")
def api_status():
    return jsonify({
        "ollama": extractor.backend_status(),
        "job": _job_public(_active_job()),
    })


@app.route("/api/sources")
def api_sources():
    return jsonify([
        {"name": k, "label": v["label"]} for k, v in searchers.SOURCE_REGISTRY.items()
    ])


@app.route("/api/settings", methods=["GET", "POST"])
def api_settings():
    if request.method == "POST":
        s = load_settings()
        data = request.get_json(force=True, silent=True) or {}
        if data.get("email") is not None:
            s["email"] = str(data["email"]).strip()
        if data.get("model") is not None:
            s["model"] = str(data["model"]).strip()
        if data.get("max_results") is not None:
            try:
                s["max_results"] = int(data["max_results"])
            except (TypeError, ValueError):
                pass
        save_settings(s)
        apply_settings(s)
        return jsonify(s)
    return jsonify(load_settings())


@app.route("/api/search", methods=["POST"])
def api_search():
    data = request.get_json(force=True, silent=True) or {}
    query = (data.get("query") or "").strip()
    sources = data.get("sources") or []
    max_results = data.get("max_results") or load_settings().get("max_results") or config.DEFAULT_MAX_RESULTS
    if not query:
        return jsonify({"error": "Enter at least one search term."}), 400
    if not sources:
        return jsonify({"error": "Choose at least one website."}), 400
    try:
        max_results = int(max_results)
    except (TypeError, ValueError):
        max_results = config.DEFAULT_MAX_RESULTS
    job, err = _start_job("search", _run_search_job, query, sources, max_results)
    if err:
        return jsonify({"error": err}), 409
    return jsonify({"job_id": job["id"]})


@app.route("/api/extract", methods=["POST"])
def api_extract():
    data = request.get_json(force=True, silent=True) or {}
    goal = (data.get("goal") or "").strip()
    columns = data.get("columns") or []
    if not goal:
        return jsonify({"error": "Describe what data you want to extract."}), 400
    job, err = _start_job("extract", _run_extract_job, goal, columns)
    if err:
        return jsonify({"error": err}), 409
    return jsonify({"job_id": job["id"]})


@app.route("/api/job/<job_id>")
def api_job(job_id):
    job = JOBS.get(job_id)
    if not job:
        return jsonify({"error": "unknown job"}), 404
    return jsonify(_job_public(job))


def _read_xlsx(path):
    p = Path(path)
    if not p.exists():
        return None
    try:
        wb = load_workbook(p, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return None
    if not rows:
        return None
    header = ["" if c is None else c for c in rows[0]]
    data = [["" if c is None else c for c in r] for r in rows[1:]]
    return {"header": header, "rows": data}


@app.route("/api/results")
def api_results():
    return jsonify({
        "search": _read_xlsx(config.SEARCH_RESULTS_XLSX),
        "extracted": _read_xlsx(config.EXTRACTED_XLSX),
        "closed": _read_xlsx(config.CLOSED_SOURCE_XLSX),
    })


@app.route("/download/<name>")
def download(name):
    mapping = {
        "search_results.xlsx": config.SEARCH_RESULTS_XLSX,
        "extracted_data.xlsx": config.EXTRACTED_XLSX,
        "closed_source_papers.xlsx": config.CLOSED_SOURCE_XLSX,
    }
    path = mapping.get(name)
    if not path or not Path(path).exists():
        abort(404)
    return send_file(str(path), as_attachment=True, download_name=name)


def _open_folder(path):
    if sys.platform == "win32":
        os.startfile(str(path))
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path)])


@app.route("/open-output", methods=["POST"])
def open_output():
    # The files already live in output/ — this just opens that folder in the
    # OS file manager so users don't have to hunt for it.
    try:
        _open_folder(config.OUTPUT_DIR)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    return jsonify({"ok": True})


if __name__ == "__main__":
    if config.APP_AUTO_OPEN:
        url = f"http://{config.APP_HOST}:{config.APP_PORT}"
        threading.Timer(1.0, lambda: webbrowser.open(url)).start()
    app.run(
        host=config.APP_HOST,
        port=config.APP_PORT,
        debug=False,
        use_reloader=False,
        threaded=True,
    )
