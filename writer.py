"""
writer.py — write the three Excel outputs using openpyxl (no pandas).

Files:
  search_results.xlsx        (overwritten each search run)
  closed_source_papers.xlsx  (appended as papers are found closed)
  extracted_data.xlsx        (appended as papers are extracted)
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    return value


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)


def _hyperlink_doi(ws, columns):
    """Make any 'doi' column clickable (https://doi.org/<doi>)."""
    header = [str(c) for c in columns]
    if "doi" not in header:
        return
    col = header.index("doi") + 1  # openpyxl columns are 1-indexed
    for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        cell = row[0]
        doi = str(cell.value or "").strip()
        if doi:
            url = doi if doi.startswith(("http://", "https://")) else f"https://doi.org/{doi}"
            cell.hyperlink = url
            cell.font = Font(color="0563C1", underline="single")


def write_xlsx(path, columns, rows):
    """Overwrite `path` with a fresh sheet: header + rows."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(columns)
    for row in rows:
        ws.append([_cell(row.get(c, "")) for c in columns])
    _style_header(ws)
    _hyperlink_doi(ws, columns)
    wb.save(path)


def append_xlsx(path, columns, rows):
    """Append rows to `path` (create it with a header first if missing)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "data"
        ws.append(columns)
    for row in rows:
        ws.append([_cell(row.get(c, "")) for c in columns])
    _style_header(ws)
    _hyperlink_doi(ws, columns)
    wb.save(path)
